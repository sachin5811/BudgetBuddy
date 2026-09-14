import io
from datetime import date

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.orm import Session

from app.models.expense import Expense
from app.models.income import Income
from app.models.user import User
from app.services.analytics_service import category_spending, summary


def generate_excel(db: Session, user: User, month: str) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    s = summary(db, user.id, month)
    ws.append(["BudgetBuddy Financial Report"])
    ws.append(["User", user.name])
    ws.append(["Email", user.email])
    ws.append(["Month", month])
    ws.append([])
    ws.append(["Metric", "Value"])
    ws.append(["Total Income", s["total_income"]])
    ws.append(["Total Expense", s["total_expense"]])
    ws.append(["Balance", s["balance"]])
    ws.append(["Total Saved", s["total_saved"]])

    ws2 = wb.create_sheet("Category Spending")
    ws2.append(["Category", "Amount"])
    for row in category_spending(db, user.id, month):
        ws2.append([row["category"], row["amount"]])

    ws3 = wb.create_sheet("Expenses")
    ws3.append(["Date", "Category", "Amount", "Description"])
    expenses = db.query(Expense).filter(Expense.user_id == user.id).all()
    for e in expenses:
        if e.date and e.date.strftime("%Y-%m") == month:
            ws3.append(
                [e.date.isoformat(), e.category, e.amount, e.description or ""]
            )

    ws4 = wb.create_sheet("Income")
    ws4.append(["Date", "Source", "Amount", "Description"])
    incomes = db.query(Income).filter(Income.user_id == user.id).all()
    for i in incomes:
        if i.date and i.date.strftime("%Y-%m") == month:
            ws4.append(
                [i.date.isoformat(), i.source, i.amount, i.description or ""]
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_pdf(db: Session, user: User, month: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title="BudgetBuddy Report")
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("BudgetBuddy Financial Report", styles["Title"]))
    elements.append(Spacer(1, 0.3 * cm))
    elements.append(
        Paragraph(f"<b>User:</b> {user.name} ({user.email})", styles["Normal"])
    )
    elements.append(Paragraph(f"<b>Month:</b> {month}", styles["Normal"]))
    elements.append(
        Paragraph(f"<b>Generated:</b> {date.today().isoformat()}", styles["Normal"])
    )
    elements.append(Spacer(1, 0.5 * cm))

    s = summary(db, user.id, month)
    summary_data = [
        ["Metric", "Value"],
        ["Total Income", f"{s['total_income']}"],
        ["Total Expense", f"{s['total_expense']}"],
        ["Balance", f"{s['balance']}"],
        ["Total Saved", f"{s['total_saved']}"],
    ]
    summary_table = Table(summary_data, colWidths=[8 * cm, 6 * cm])
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#059669")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("PADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(summary_table)
    elements.append(Spacer(1, 0.6 * cm))

    elements.append(Paragraph("Category-wise Spending", styles["Heading2"]))
    cat_data = [["Category", "Amount"]]
    for row in category_spending(db, user.id, month):
        cat_data.append([row["category"], f"{row['amount']}"])
    if len(cat_data) == 1:
        cat_data.append(["No expenses", "0"])
    cat_table = Table(cat_data, colWidths=[8 * cm, 6 * cm])
    cat_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                ("PADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(cat_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
